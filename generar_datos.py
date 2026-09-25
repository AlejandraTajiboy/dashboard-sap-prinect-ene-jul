"""
Lee el Excel SAP + Prinect (solo lectura, nunca lo modifica) y genera data.json para el dashboard.

Uso:
    python generar_datos.py                      # usa la ruta por defecto (ENE-JUL)
    python generar_datos.py "C:\\ruta\\otro.xlsx"  # usa otro archivo

Fuente principal: hoja «desviacion vs equivalente» (en agosto se llamaba «Desviacion_Cantidad_vs_Equiv»).
Los valores de cada OP (equivalente, recibida, desviaciones, valor monetario, costos, variación y desviación
Prinect) se toman tal como los calcula Excel.
Hojas de apoyo:
  - general : pliegos de prensa SAP y Prinect por OP («PLIEGOS PRENSA SAP» y «PLIEGOS PRENSA PRINECT»),
              y un recálculo independiente del equivalente y de la cantidad recibida para validar la hoja base
  - prinect : (en agosto «Hoja1») para saber si la OP existe en Prinect y cuántas corridas y cortes tiene

Los encabezados se buscan sin importar mayúsculas, espacios ni acentos, así que sirve para los dos formatos.

Cada OP recibe un estado por sección. «Con datos» entra en la gráfica y en los totales. Los demás estados
(«Sin rollo», «Sin entrada», «No está en Prinect», etc.) dicen por qué esa OP no se puede comparar.
El filtro por mes se aplica en la página; aquí cada OP lleva su mes (de «Fecha de contabilización»).
"""
import json
import math
import shutil
import sys
import tempfile
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

RUTA_POR_DEFECTO = r"C:\Users\dtajiboy\Documents\Dosificadacion\Movimiento\ENE-JUL\ENE-JUL_MOVIMIENTO_SAPPRINECT.xlsx"
SALIDA_JSON = Path(__file__).with_name("data.json")
ANIO, MESES_PERIODO = 2026, range(1, 8)  # enero–julio 2026
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

MM_A_PULG = 0.0393701
M_A_PULG = 39.3701
UM_ROLLO = {"metros", "metro(s)"}
UM_PLIEGO = {"hoja(s)", "pliegos", "pliego", "pliego(s)"}
CON_DATOS = "Con datos"

# Cada columna acepta varios nombres (agosto / ENE-JUL). Se comparan con norm().
HOJA_DESV = ("desviacion_cantidad_vs_equiv", "desviacion vs equivalente")
HOJA_PRINECT = ("prinect", "hoja1")
COL_DESV = {
    "op": ("no.pedido (op)",),
    "fecha": ("fecha de contabilización",),
    "articulo": ("número de artículo",),
    "descripcion": ("descripción",),
    "equivalente": ("cantidad equivalente (rollo -> pliegos)", "cantidad equivalente"),
    "recibida": ("cantidad real recibida (entrada)", "cantidad real recibida"),
    "desv_abs": ("desviación absoluta",),
    "desv_pct": ("desviación %",),
    "estado_excel": ("estado",),
    "valor": ("valor monetario",),
    "costo_ideal": ("costo unidad ideal",),
    "costo_sistema": ("costo unidad sistema",),
    "costo_var": ("variacion",),
    "costo_pct": ("% variacion", "variacion %"),
    "prinect_desv": ("des prinect/sap", "desviacion prinect/sap"),
    "prinect_pct": ("des % prinect/sap", "desviacion % prinect/sap"),
}
COL_GENERAL = {
    "comentarios": ("comentarios",),
    "cantidad": ("cantidad",),
    "um": ("nombre de unidad de medida",),
    "altura": ("altura/largo",),
    "op": ("no.pedido",),
    "cortes": ("cortes",),
    "pliego_sap": ("pliego prensa sap", "pliegos prensa sap"),
    "pliego_prinect": ("pliego prensa prinect", "pliegos prensa prinect"),
}
COL_PRINECT = {"op": ("op",), "titulo": ("titulo",), "cantidad": ("cantidad", "total ciclos"), "cortes": ("cortes", "corte")}


# ---------------------------------------------------------------- utilidades
def norm(s):
    """Minúsculas, sin acentos, sin espacios ni guiones bajos: «DESVIACION %» == «Desviación %»."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return "".join(ch for ch in s.lower() if not ch.isspace() and ch != "_")


def num(v):
    """Número o None. Los errores de Excel (#DIV/0!, #N/A) y el texto cuentan como None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def op_id(v):
    n = num(v)
    if n is None:
        return None
    return str(int(n)) if n == int(n) else str(n)


def fecha_iso(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, datetime) else None


def r(v, d=4):
    return round(v, d) if v is not None else None


def pct(v):
    n = num(v)
    return round(n * 100, 4) if n is not None else None


def es_error(v):
    return isinstance(v, str) and v.startswith("#")


def rango_de(pct_abs):
    """0-10, 11-20, ... 91-100, >100 (en porcentaje, valor absoluto)."""
    if pct_abs > 100:
        return ">100%"
    b = max(1, math.ceil(round(pct_abs, 9) / 10))
    return "0–10%" if b == 1 else f"{(b - 1) * 10 + 1}–{b * 10}%"


RANGOS = ["0–10%"] + [f"{(b - 1) * 10 + 1}–{b * 10}%" for b in range(2, 11)] + [">100%"]


def leer_tabla(ws, columnas, requeridas=()):
    filas = ws.iter_rows(values_only=True)
    encabezado = [norm(h) for h in next(filas)]
    idx = {}
    for k, nombres in columnas.items():
        i = next((encabezado.index(norm(n)) for n in nombres if norm(n) in encabezado), None)
        if i is not None:
            idx[k] = i
        elif k in requeridas:
            sys.exit(f"La hoja «{ws.title}» no tiene la columna «{nombres[-1]}».")
    out = []
    for f in filas:
        if f is None or all(c is None for c in f):
            continue
        out.append({k: (f[i] if i < len(f) else None) for k, i in idx.items()})
    return out


def leer_libro(ruta):
    # Se trabaja sobre una copia temporal: el original no se toca aunque esté abierto en Excel.
    with tempfile.TemporaryDirectory() as tmp:
        copia = Path(tmp) / "copia.xlsx"
        shutil.copy2(ruta, copia)
        wb = openpyxl.load_workbook(copia, read_only=True, data_only=True)
        hojas = {norm(h): h for h in wb.sheetnames}
        buscar = lambda nombres: next((hojas[norm(n)] for n in nombres if norm(n) in hojas), None)
        h_desv, h_pri = buscar(HOJA_DESV), buscar(HOJA_PRINECT)
        if not h_desv:
            sys.exit(f"El archivo no tiene la hoja «{HOJA_DESV[-1]}».")
        desv = leer_tabla(wb[h_desv], COL_DESV, ("op", "equivalente", "recibida"))
        general = leer_tabla(wb[hojas["general"]], COL_GENERAL, ("op", "um", "cantidad")) if "general" in hojas else []
        prinect = leer_tabla(wb[h_pri], COL_PRINECT, ("op",)) if h_pri else []
        wb.close()
    return h_desv, h_pri or "prinect", desv, general, prinect


# ---------------------------------------------------------------- estados por sección
def estado_cantidad(o):
    if not o["recibida"]:
        return "Sin entrada"
    if o["entrada_metros"]:
        return "Entrada en metros"
    if not o["equivalente"]:
        return "Sin rollo"
    return CON_DATOS if o["desv_pct"] is not None else "Sin cálculo"


def estado_costo(o):
    if not o["recibida"]:
        return "Sin entrada"
    if o["entrada_metros"]:
        return "Entrada en metros"
    if not o["equivalente"]:
        return "Sin rollo"
    return CON_DATOS if o["costo_pct"] is not None else "Sin cálculo"


def estado_prinect(o, en_prinect, con_cortes):
    if not o["recibida"]:
        return "Sin entrada SAP"
    if o["entrada_metros"]:
        return "Entrada en metros"
    if not en_prinect:
        return "No está en Prinect"
    if not con_cortes:
        return "Sin cortes en Prinect"
    return CON_DATOS if o["prinect_pct"] is not None else "Sin cálculo"


def lista(ops, n=12):
    ids = [o["op"] if isinstance(o, dict) else o for o in ops]
    return ", ".join(ids[:n]) + (f" y {len(ids) - n} más" if len(ids) > n else "")


# ---------------------------------------------------------------- principal
def construir(ruta):
    hoja_base, hoja_prinect, desv, general, prinect_filas = leer_libro(ruta)

    es_total = lambda d: str(d.get("descripcion") or "").strip().upper() == "TOTAL"
    total_excel = next((d for d in desv if es_total(d)), None)
    filas = [d for d in desv if op_id(d["op"]) and not es_total(d)]

    # Apoyo desde «general»: pliegos de prensa por OP y recálculo del equivalente / recibida
    g_sap, g_pri, g_cortes = defaultdict(float), defaultdict(float), {}
    n_entradas, metros, recibida_g, alturas = Counter(), defaultdict(float), defaultdict(float), defaultdict(list)
    ent_metros, ent_pliegos, salidas = set(), set(), set()
    for g in general:
        op = op_id(g["op"])
        um = str(g["um"] or "").strip().lower()
        entrada = str(g.get("comentarios") or "").strip().lower().startswith("entrada") if "comentarios" in g else um in UM_PLIEGO
        if um in UM_PLIEGO:
            ent_pliegos.add(op)
            n_entradas[op] += 1
            g_sap[op] += num(g.get("pliego_sap")) or 0
            g_pri[op] += num(g.get("pliego_prinect")) or 0
            g_cortes.setdefault(op, num(g.get("cortes")))
            recibida_g[op] += num(g["cantidad"]) or 0
            if (num(g["altura"]) or 0) > 0:
                alturas[op].append(num(g["altura"]) * MM_A_PULG)
        elif um in UM_ROLLO and entrada:
            ent_metros.add(op)
        elif um in UM_ROLLO:
            salidas.add(op)
            metros[op] += num(g["cantidad"]) or 0

    # Apoyo desde «prinect»
    p_by = defaultdict(list)
    p_raras = []
    for p in prinect_filas:
        pid = op_id(p["op"])
        if pid:
            p_by[pid].append(p)
        elif p["op"] is not None:
            p_raras.append(str(p["op"]).strip())

    ops, errores = [], Counter()
    for d in filas:
        op = op_id(d["op"])
        ps = p_by.get(op, [])
        cortes_corrida = [num(p["cortes"]) or 0 for p in ps]
        desc = d["descripcion"]
        for k in ("desv_pct", "costo_ideal", "costo_pct", "prinect_pct"):
            if es_error(d.get(k)):
                errores[k] += 1
        fe = d["fecha"] if isinstance(d["fecha"], datetime) else None
        o = {
            "op": op,
            "fecha": fecha_iso(fe),
            "mes": fe.month if fe and fe.year == ANIO else 0,
            "articulo": op_id(d["articulo"]) if num(d["articulo"]) else None,
            "descripcion": desc if isinstance(desc, str) and desc.strip() not in ("", "0") else None,
            "estado_excel": d.get("estado_excel"),
            "equivalente": r(num(d["equivalente"])),
            "recibida": r(num(d["recibida"])),
            "desv_abs": r(num(d["desv_abs"])),
            "desv_pct": pct(d["desv_pct"]),
            "valor": r(num(d.get("valor")), 2),
            "costo_ideal": r(num(d.get("costo_ideal")), 6),
            "costo_sistema": r(num(d.get("costo_sistema")), 6),
            "costo_var": r(num(d.get("costo_var")), 6),
            "costo_pct": pct(d.get("costo_pct")),
            "prinect_desv": r(num(d.get("prinect_desv")), 2),
            "prinect_pct": pct(d.get("prinect_pct")),
            "pliegos_sap": r(g_sap[op], 2) if op in g_sap else None,
            "pliegos_prinect": r(g_pri[op], 2) if op in g_pri else None,
            "cortes": g_cortes.get(op),
            "cortes_corrida": max(cortes_corrida) if cortes_corrida else None,
            "filas_entrada": n_entradas.get(op, 0),
            "entrada_metros": op in ent_metros and op not in ent_pliegos,
            "sin_salida": op not in salidas,
            "prinect_corridas": len(ps),
            "prinect_titulo": next((str(p["titulo"]) for p in ps if p.get("titulo")), None),
        }
        o["est_cantidad"] = estado_cantidad(o)
        o["est_costo"] = estado_costo(o)
        o["est_prinect"] = estado_prinect(o, bool(ps), any(cortes_corrida))
        o["pct_div0"] = es_error(d["desv_pct"])
        # Sin datos: no se muestran porcentajes que en Excel son #DIV/0! o 0 sin sentido
        if o["est_cantidad"] != CON_DATOS:
            o["desv_pct"] = o["desv_abs"] = None
        if o["est_costo"] != CON_DATOS:
            o["costo_ideal"] = o["costo_var"] = o["costo_pct"] = None
        if o["est_prinect"] != CON_DATOS:
            o["prinect_pct"] = None
        o["rango"] = rango_de(abs(o["desv_pct"])) if o["est_cantidad"] == CON_DATOS else None
        o["costo_rango"] = rango_de(abs(o["costo_pct"])) if o["est_costo"] == CON_DATOS else None
        o["prinect_rango"] = rango_de(abs(o["prinect_pct"])) if o["est_prinect"] == CON_DATOS else None
        ops.append(o)

    periodo = [o for o in ops if o["mes"] in MESES_PERIODO]
    fuera = [o for o in ops if o["mes"] not in MESES_PERIODO]
    qc = [o for o in periodo if o["est_cantidad"] == CON_DATOS]
    cc = [o for o in periodo if o["est_costo"] == CON_DATOS]
    pc = [o for o in periodo if o["est_prinect"] == CON_DATOS]
    tot_eq = sum(o["equivalente"] for o in qc)
    tot_rec = sum(o["recibida"] for o in qc)

    # ------------------------------------------------ validación: equivalente y recibida recalculados desde «general»
    q_all = [o for o in ops if o["est_cantidad"] == CON_DATOS]
    dif, sin_recalc = [], 0
    for o in q_all:
        alt = alturas.get(o["op"])
        if not (alt and metros.get(o["op"])):
            sin_recalc += 1
            continue
        eq = metros[o["op"]] * M_A_PULG / (sum(alt) / len(alt))
        if abs(eq - o["equivalente"]) > 0.01 or abs(recibida_g[o["op"]] - o["recibida"]) > 0.01:
            dif.append(o["op"])
    validacion = {"ops_revisadas": len(q_all) - sin_recalc, "ops_con_diferencia": len(dif), "ops_diferencia": dif[:20]}

    # ------------------------------------------------ revisión de la lógica de ESTE archivo
    notas = []
    if fuera:
        c_mes = Counter(o["mes"] for o in fuera)
        det = ", ".join(f"{c_mes[m]} de {MESES[m - 1].lower()}" if m else f"{c_mes[m]} sin fecha" for m in sorted(c_mes))
        sin_e = sum(1 for o in fuera if o["est_cantidad"] == "Sin entrada")
        notas.append(
            f"El archivo se llama ENE-JUL, pero {len(fuera)} OP tienen «Fecha de contabilización» fuera de enero–julio ({det}); "
            f"{sin_e} de ellas aún no tienen entrada. No entran en los totales del periodo. Se pueden ver con el filtro de mes «Fuera del periodo»."
        )
    sin_rollo = [o for o in ops if o["est_cantidad"] == "Sin rollo"]
    en_metros = [o for o in ops if o["est_cantidad"] == "Entrada en metros"]
    sin_ent = [o for o in ops if o["est_cantidad"] == "Sin entrada"]
    if total_excel is None:
        all_e = sum(o["equivalente"] or 0 for o in ops)
        all_r = sum(o["recibida"] or 0 for o in ops)
        notas.append(
            f"La hoja «{hoja_base}» no tiene fila TOTAL. Si se suman las columnas completas, la desviación da "
            f"{(all_r - all_e) / all_e * 100:+.2f}% (recibida {all_r:,.0f} vs. equivalente {all_e:,.0f}), porque se suman los pliegos de las "
            f"{len(sin_rollo) + len(en_metros)} OP con equivalente 0 (sin rollo o entrada en metros) y los metros como si fueran pliegos. "
            f"Con solo las OP con datos de enero–julio, la desviación es {(tot_rec - tot_eq) / tot_eq * 100 if tot_eq else 0:+.2f}%."
        )
    else:
        te, tr = num(total_excel["equivalente"]), num(total_excel["recibida"])
        notas.append(
            f"El TOTAL de la hoja da {(tr - te) / te * 100:+.2f}% porque suma las OP sin rollo (equivalente 0). Con solo las OP con datos, "
            f"la desviación es {(tot_rec - tot_eq) / tot_eq * 100 if tot_eq else 0:+.2f}%."
        )
    mal_estado = [o for o in sin_rollo + en_metros if str(o["estado_excel"]).strip() == "Comparado"]
    if mal_estado:
        div0 = sum(1 for o in mal_estado if o["pct_div0"])
        notas.append(
            f"«ESTADO» marca como «Comparado» a {len(mal_estado)} OP con equivalente 0 ({div0} con «DESVIACION %» = #DIV/0!). "
            "La condición «Sin equivalente» nunca se cumple porque SUMAR.SI.CONJUNTO devuelve 0, no vacío. Además, «DESVIACION ABSOLUTA» "
            "queda igual a toda la cantidad recibida. En el dashboard aparecen como «Sin rollo» o «Entrada en metros»."
        )
    if en_metros:
        notas.append(
            f"{len(en_metros)} OP tienen la entrada registrada en METROS y no en pliegos ({lista(en_metros)}). «cantidad real recibida» suma "
            "esos metros como si fueran pliegos, y el equivalente sale 0 porque PROMEDIO.SI.CONJUNTO no encuentra altura de pliego: «general» "
            "devuelve el texto «Sin entrada con esa OP» aunque sí hay entrada. En el dashboard aparecen como «Entrada en metros»."
        )
    sr = [o for o in sin_rollo if o["sin_salida"]]
    if sin_rollo:
        notas.append(
            f"{len(sin_rollo)} OP tienen entrada de pliegos pero no salida de rollo ({lista(sin_rollo)})"
            + (f"; {len(sr)} no aparecen en la hoja «salida»" if sr else "")
            + ". Su equivalente es 0 y su desviación % da #DIV/0!. En el dashboard aparecen como «Sin rollo»."
        )
    if sin_ent:
        notas.append(
            f"{len(sin_ent)} OP tienen salida de rollo pero no entrada. Su equivalente sale en 0 porque «general» devuelve el texto "
            "«Sin entrada con esa OP» y SUMAR.SI.CONJUNTO lo ignora. Excel sí las marca «Sin entrada (hoja/pliegos)». En el dashboard aparecen como «Sin entrada»."
        )
    if errores["costo_ideal"]:
        notas.append(
            f"«Costo Unidad Ideal» (=Valor Monetario ÷ Cantidad equivalente) da #DIV/0! en {errores['costo_ideal']} OP con equivalente 0, "
            "y el error pasa a «VARIACION» y «VARIACION %». Convendría envolverlo en SI.ERROR o marcar esas OP como sin datos."
        )
    multi_c = [o for o in ops if o["filas_entrada"] > 1 and o["est_costo"] == CON_DATOS]
    if multi_c:
        ej = max(multi_c, key=lambda o: abs(o["costo_pct"]))
        notas.append(
            f"«Costo unidad Sistema» es SUMAR.SI.CONJUNTO del costo unitario de cada fila de entrada. En las {len(multi_c)} OP con datos que tienen "
            f"varias entradas ({lista(multi_c)}) el costo queda sumado en vez de promediado y la variación sale inflada "
            f"(OP {ej['op']}: {ej['costo_pct']:+.2f}%). Se corregiría con SUMAR.SI.CONJUNTO del Total de entrada ÷ cantidad recibida."
        )
    cortes_sumados = [o for o in ops if o["prinect_corridas"] > 1 and o["cortes"] and o["cortes_corrida"] and o["cortes"] > o["cortes_corrida"]]
    if cortes_sumados:
        ej = cortes_sumados[0]
        notas.append(
            f"«CORTES» en «general» es SUMAR.SI.CONJUNTO sobre «{hoja_prinect}». Si la OP tiene varias corridas en Prinect, suma los cortes de "
            f"todas: en {len(cortes_sumados)} OP ({lista(cortes_sumados)}) los cortes y los «PLIEGOS PRENSA SAP» salen multiplicados. "
            f"Ejemplo: OP {ej['op']} tiene {ej['prinect_corridas']} corridas de {ej['cortes_corrida']:g} corte(s) y queda con {ej['cortes']:g}. "
            "Para los cortes conviene MAX.SI.CONJUNTO. Para los ciclos de Prinect, SUMAR.SI.CONJUNTO sí es correcto."
        )
    multi_p = [o for o in ops if o["filas_entrada"] > 1 and o["pliegos_prinect"]]
    if multi_p:
        notas.append(
            f"«PLIEGOS PRENSA PRINECT» se trae completo en cada fila de entrada. En las {len(multi_p)} OP con varias entradas "
            f"({lista(multi_p)}), «DESVIACION PRINECT/SAP» cuenta Prinect dos o más veces y «DESVIACION % PRINECT/SAP» suma los "
            "porcentajes de cada fila."
        )
    no_p = [o for o in ops if o["est_prinect"] == "No está en Prinect"]
    sin_c = [o for o in ops if o["est_prinect"] == "Sin cortes en Prinect"]
    if errores["prinect_pct"]:
        notas.append(
            f"«DESVIACION % PRINECT/SAP» da #DIV/0! en {errores['prinect_pct']} OP: {len(no_p)} no existen en «{hoja_prinect}» y {len(sin_c)} tienen "
            "CORTE = 0 en Prinect (los demás son OP con entrada en metros). En el dashboard aparecen como «No está en Prinect» y «Sin cortes en Prinect»."
        )
    notas.append(
        "«PLIEGOS PRENSA PRINECT» usa «TOTAL CICLOS» (ciclos buenos + malos). Si la comparación con SAP debe ser solo con producción buena, "
        "habría que usar «CICLOS BUENOS»."
    )
    largo = Counter(len(o["op"]) for o in ops).most_common(1)[0][0] if ops else 0
    raras = [o["op"] for o in ops if len(o["op"]) != largo]
    if raras or p_raras:
        txt = f"OP con formato distinto al resto (posible error de captura): {', '.join(raras)}." if raras else ""
        if p_raras:
            txt += (" " if txt else "") + (
                f"En «{hoja_prinect}» hay {len(p_raras)} filas con OP de texto ({', '.join(sorted(set(p_raras)))}), que son calibraciones y pruebas; "
                "no cruzan con ninguna OP de SAP."
            )
        notas.append(txt)
    notas.append(
        "«Fecha de contabilización», «Número de artículo» y «Descripción» se traen con BUSCARX, que devuelve la primera fila de la OP en "
        "«general» (la primera salida, no necesariamente la más antigua). El mes del filtro sale de esa fecha."
    )

    fechas = [o["fecha"] for o in periodo if o["fecha"]]
    meses = [{"mes": m, "nombre": MESES[m - 1], "ops": sum(1 for o in ops if o["mes"] == m), "periodo": m in MESES_PERIODO}
             for m in sorted({o["mes"] for o in ops if o["mes"]})]
    for o in ops:
        for k in ("pct_div0", "sin_salida", "entrada_metros"):
            o.pop(k)
    return {
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "archivo": Path(ruta).name,
        "hoja_base": hoja_base,
        "titulo_periodo": f"enero–julio {ANIO}",
        "periodo": {"desde": min(fechas) if fechas else None, "hasta": max(fechas) if fechas else None},
        "meses": meses,
        "rangos": RANGOS,
        "ops": ops,
        "notas": notas,
        "validacion": validacion,
        "_consola": {
            "ops": len(ops), "periodo": len(periodo), "fuera": len(fuera),
            "cantidad": (len(qc), (tot_rec - tot_eq) / tot_eq * 100 if tot_eq else 0, Counter(o["est_cantidad"] for o in periodo)),
            "costo": (len(cc), Counter(o["est_costo"] for o in periodo)),
            "prinect": (len(pc), Counter(o["est_prinect"] for o in periodo)),
            "sin_datos": sum(1 for o in periodo if o["est_cantidad"] != CON_DATOS or o["est_costo"] != CON_DATOS or o["est_prinect"] != CON_DATOS),
        },
    }


def main():
    ruta = sys.argv[1] if len(sys.argv) > 1 else RUTA_POR_DEFECTO
    if not Path(ruta).exists():
        sys.exit(f"No se encontró el archivo: {ruta}")
    data = construir(ruta)
    c = data.pop("_consola")
    SALIDA_JSON.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"OK -> {SALIDA_JSON}  (base: hoja {data['hoja_base']})")
    print(f"  OP en la hoja: {c['ops']} | enero-julio: {c['periodo']} | fuera del periodo: {c['fuera']}")
    print(f"  Enero-julio, con algún dato faltante: {c['sin_datos']}")
    print(f"  Cantidad: {c['cantidad'][0]} OP con datos | {c['cantidad'][1]:+.2f}% | {dict(c['cantidad'][2])}")
    print(f"  Costo:    {c['costo'][0]} OP con datos | {dict(c['costo'][1])}")
    print(f"  Prinect:  {c['prinect'][0]} OP con datos | {dict(c['prinect'][1])}")
    v = data["validacion"]
    print(f"  Validación (recalculado desde «general»): {v['ops_revisadas']} OP, {v['ops_con_diferencia']} con diferencia")


if __name__ == "__main__":
    main()
